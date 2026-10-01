# utils/excel_generator.py
"""
Разбор JSON-ответа ИИ в pandas.DataFrame и генерация xlsx по JSON-шаблону
(templates/excel_template.json, загружается в config.EXCEL_TEMPLATE).

Код не знает конкретных полей: колонки, заголовки, группы, формулы,
верхние и нижние служебные строки берутся из шаблона. Чтобы добавить
или изменить колонку, правьте шаблон (и промпт, если поле приходит от ИИ).

Формат шаблона
--------------
columns[] — список колонок слева направо:
    key           — внутреннее имя (ключ в JSON ответа ИИ и колонка DataFrame)
    label         — заголовок в Excel
    source        — "ai" (значение из ответа модели),
                    "manual" (пустая, заполняет пользователь),
                    "formula" (формула Excel, поле "formula")
    group         — (необяз.) общий заголовок над соседними колонками
    type          — (необяз.) "list": значение — массив, каждый элемент
                    выводится в отдельной строке
    item_level    — (необяз., по умолчанию true) ячейки колонки объединяются
                    по вертикали, если у позиции несколько строк
    wrap          — (необяз.) перенос по словам
    width         — (необяз.) ширина колонки; иначе подбирается автоматически
    number_format — (необяз.) числовой формат Excel
    formula       — для source="formula"

top_rows[]    — строки НАД шапкой: id, label, input_column
                (под подписью объединяются ячейки до input_column,
                сама ячейка ввода остаётся пустой)
footer_rows[] — строки ПОД таблицей (итоги). В одной строке может быть
                несколько итоговых ячеек:
                    label        — подпись строки
                    label_column — (необяз.) колонка подписи, по умолчанию первая
                    cells[]      — итоговые ячейки: column, formula,
                                   number_format (необяз.)
                Подпись объединяется от label_column до первой итоговой ячейки.
                Старый формат value_column + formula читается как одна ячейка.

Плейсхолдеры в формулах:
    {key}        — ячейка колонки key в текущей строке позиции
    {key:range}  — диапазон данных колонки key (первая—последняя строка)
    {@id}        — абсолютный адрес ячейки ввода из top_rows

Формулы пишутся с английскими именами функций и запятой между
аргументами: так они хранятся в xlsx, Excel сам покажет их в локали
пользователя (ЕСЛИ, И, ЕЧИСЛО, ;).
"""
from __future__ import annotations

import io
import json
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from config import EXCEL_TEMPLATE


class ResponseParsingError(Exception):
    """Ошибка разбора ответа ИИ в табличный вид."""


class TemplateError(Exception):
    """Ошибка в JSON-шаблоне Excel."""


_PLACEHOLDER_RE = re.compile(r"\{(@?)([A-Za-z0-9_]+)(?::(range))?\}")
_VALID_SOURCES = {"ai", "manual", "formula"}


# ---------------------------------------------------------------------------
# Работа с шаблоном
# ---------------------------------------------------------------------------

def _is_list_column(col: dict) -> bool:
    return col.get("type") == "list"


def _is_item_level(col: dict) -> bool:
    return col.get("item_level", True)


def validate_template(template: dict) -> None:
    """Проверяет шаблон и бросает TemplateError с понятным текстом."""
    columns = template.get("columns")
    if not isinstance(columns, list) or not columns:
        raise TemplateError("В шаблоне нет непустого списка 'columns'.")

    keys: list[str] = []
    for col in columns:
        key = col.get("key")
        if not key:
            raise TemplateError(f"У колонки нет 'key': {col}")
        if key in keys:
            raise TemplateError(f"Ключ колонки повторяется: {key}")
        keys.append(key)
        if "label" not in col:
            raise TemplateError(f"У колонки '{key}' нет 'label'.")
        source = col.get("source")
        if source not in _VALID_SOURCES:
            raise TemplateError(
                f"У колонки '{key}' неверный source: {source!r} "
                f"(допустимо: {sorted(_VALID_SOURCES)})."
            )
        if source == "formula" and not col.get("formula"):
            raise TemplateError(f"У формульной колонки '{key}' нет 'formula'.")

    top_ids = []
    for row in template.get("top_rows", []):
        if row.get("input_column") not in keys:
            raise TemplateError(f"top_rows: неизвестная input_column в {row}")
        if not row.get("id"):
            raise TemplateError(f"top_rows: у строки нет 'id': {row}")
        top_ids.append(row["id"])

    for row in template.get("footer_rows", []):
        label_key = row.get("label_column", keys[0])
        if label_key not in keys:
            raise TemplateError(f"footer_rows: неизвестная label_column в {row}")
        cells = _footer_cells(row)
        if not cells:
            raise TemplateError(f"footer_rows: у строки нет итоговых ячеек: {row}")
        used: set[str] = set()
        for cell in cells:
            column = cell.get("column")
            if column not in keys:
                raise TemplateError(f"footer_rows: неизвестная колонка {column!r} в {row}")
            if column in used:
                raise TemplateError(f"footer_rows: колонка '{column}' повторяется в одной строке: {row}")
            used.add(column)
            if keys.index(column) <= keys.index(label_key):
                raise TemplateError(
                    f"footer_rows: итоговая ячейка '{column}' должна быть правее колонки подписи '{label_key}'."
                )
            if not cell.get("formula"):
                raise TemplateError(f"footer_rows: у ячейки '{column}' нет 'formula': {row}")
            for at, name, range_flag in _PLACEHOLDER_RE.findall(cell["formula"]):
                if not at and not range_flag:
                    raise TemplateError(
                        f"footer_rows: в формуле итога {{{name}}} нужно писать как {{{name}:range}}: {cell['formula']}"
                    )

    # Проверяем, что все плейсхолдеры формул ссылаются на существующее.
    formulas = [c["formula"] for c in columns if c.get("source") == "formula"]
    for r in template.get("footer_rows", []):
        formulas += [cell["formula"] for cell in _footer_cells(r)]
    for formula in formulas:
        for at, name, _range in _PLACEHOLDER_RE.findall(formula):
            if at and name not in top_ids:
                raise TemplateError(f"В формуле неизвестная ячейка ввода {{@{name}}}: {formula}")
            if not at and name not in keys:
                raise TemplateError(f"В формуле неизвестная колонка {{{name}}}: {formula}")


def _render_formula(
    formula: str,
    col_letters: dict[str, str],
    top_cells: dict[str, str],
    row: int | None,
    first_row: int,
    last_row: int,
) -> str:
    """Подставляет адреса ячеек вместо плейсхолдеров."""

    def repl(match: re.Match) -> str:
        at, name, range_flag = match.group(1), match.group(2), match.group(3)
        if at:
            return top_cells[name]
        letter = col_letters[name]
        if range_flag:
            return f"{letter}{first_row}:{letter}{last_row}"
        if row is None:
            raise TemplateError(
                f"Плейсхолдер {{{name}}} без ':range' нельзя использовать в строке итогов."
            )
        return f"{letter}{row}"

    return _PLACEHOLDER_RE.sub(repl, formula)


# ---------------------------------------------------------------------------
# Вспомогательные функции разбора
# ---------------------------------------------------------------------------

def _is_empty(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _extract_json_block(raw_text: str) -> str:
    """Вырезает JSON-объект из ответа, даже если он обёрнут в ```json ... ```."""
    match = re.search(r"\{.*\}", raw_text, flags=re.DOTALL)
    if not match:
        raise ResponseParsingError("В ответе модели не найден JSON-объект.")
    return match.group(0)


def _normalize_list(raw_value) -> list[str]:
    """Приводит значение list-колонки к списку непустых строк."""
    if raw_value is None:
        return []
    if isinstance(raw_value, list):
        return [str(v) for v in raw_value if v is not None and str(v).strip() != ""]
    text = str(raw_value).strip()
    return [text] if text else []


def _find_item_ranges(df: pd.DataFrame, columns: list[dict]) -> list[tuple[int, int]]:
    """
    Определяет позиции по самому DataFrame (0-based, включительно).
    Строка — продолжение предыдущей позиции, если в ней нет ни одного
    значения скалярных ai-колонок, но есть значение list-колонки.
    Работает и после ручного редактирования в st.data_editor.
    """
    scalar_keys = [c["key"] for c in columns if c["source"] == "ai" and not _is_list_column(c)]
    list_keys = [c["key"] for c in columns if _is_list_column(c)]

    starts: list[int] = []
    for idx in range(len(df)):
        row = df.iloc[idx]
        has_scalar = any(not _is_empty(row[k]) for k in scalar_keys)
        has_list = any(not _is_empty(row[k]) for k in list_keys)
        if idx == 0 or has_scalar or not has_list:
            starts.append(idx)

    ranges = []
    for i, start in enumerate(starts):
        end = (starts[i + 1] - 1) if i + 1 < len(starts) else len(df) - 1
        ranges.append((start, end))
    return ranges


def _to_excel_value(value):
    """Приводит значение ячейки DataFrame к типу, понятному openpyxl."""
    if _is_empty(value):
        return None
    if hasattr(value, "item"):  # numpy-скаляры
        value = value.item()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


# ---------------------------------------------------------------------------
# Ответ ИИ -> DataFrame
# ---------------------------------------------------------------------------

def parse_ai_response_to_df(raw_response: str, template: dict | None = None) -> pd.DataFrame:
    """
    Преобразует сырой ответ ИИ в DataFrame с колонками из шаблона.

    Ожидаемые форматы ответа:
      - {"items": [...]}          — успешный разбор
      - {"error": "текст ошибки"} — модель не смогла прочитать файл

    Колонки type="list" раскладываются по строкам (элемент списка — строка).
    Значения остальных колонок записываются только в первую строку позиции.
    Колонки не от ИИ (manual/formula) остаются пустыми.

    Raises:
        ResponseParsingError: невалидный JSON, нет 'items' или ошибка модели.
    """
    template = template or EXCEL_TEMPLATE
    columns = template["columns"]
    keys = [c["key"] for c in columns]

    json_block = _extract_json_block(raw_response)
    try:
        data = json.loads(json_block)
    except json.JSONDecodeError as exc:
        raise ResponseParsingError(f"Ответ модели — невалидный JSON: {exc}") from exc

    if not isinstance(data, dict):
        raise ResponseParsingError("Ответ модели должен быть JSON-объектом.")

    if "error" in data:
        error_text = data.get("error") or "Модель сообщила об ошибке анализа файла."
        raise ResponseParsingError(str(error_text))

    items = data.get("items")
    if not isinstance(items, list) or not items:
        raise ResponseParsingError("В JSON отсутствует непустой список 'items'.")

    rows: list[dict] = []
    for item in items:
        if not isinstance(item, dict):
            continue

        lists = {c["key"]: _normalize_list(item.get(c["key"])) for c in columns if _is_list_column(c)}
        row_count = max([len(v) for v in lists.values()] + [1])

        for i in range(row_count):
            row: dict = {}
            for col in columns:
                key = col["key"]
                if _is_list_column(col):
                    row[key] = lists[key][i] if i < len(lists[key]) else None
                elif i == 0 and col["source"] == "ai":
                    row[key] = item.get(key)
                else:
                    row[key] = None
            rows.append(row)

    return pd.DataFrame(rows, columns=keys)


# ---------------------------------------------------------------------------
# DataFrame -> xlsx
# ---------------------------------------------------------------------------

def _merge_item_cells(ws, columns: list[dict], ranges: list[tuple[int, int]], data_start: int) -> None:
    """Объединяет по вертикали ячейки item_level-колонок у позиций с несколькими строками."""
    item_level_idx = [i for i, c in enumerate(columns, start=1) if _is_item_level(c)]
    for start, end in ranges:
        if end <= start:
            continue
        for col_idx in item_level_idx:
            ws.merge_cells(
                start_row=data_start + start, end_row=data_start + end,
                start_column=col_idx, end_column=col_idx,
            )


def _write_header(ws, columns: list[dict], header_row: int) -> None:
    """Двухстрочная шапка: группы объединяются по горизонтали, остальные — по вертикали."""
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    idx = 1
    while idx <= len(columns):
        col = columns[idx - 1]
        group = col.get("group")
        if group:
            end = idx
            while end + 1 <= len(columns) and columns[end].get("group") == group:
                end += 1
            ws.cell(row=header_row, column=idx).value = group
            if end > idx:
                ws.merge_cells(start_row=header_row, end_row=header_row, start_column=idx, end_column=end)
            for j in range(idx, end + 1):
                ws.cell(row=header_row + 1, column=j).value = columns[j - 1]["label"]
            idx = end + 1
        else:
            ws.cell(row=header_row, column=idx).value = col["label"]
            ws.merge_cells(start_row=header_row, end_row=header_row + 1, start_column=idx, end_column=idx)
            idx += 1

    for r in (header_row, header_row + 1):
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = header_font
            cell.alignment = header_align


def _write_label_row(ws, row: int, label: str, value_col_idx: int) -> None:
    """Подпись, прижатая вправо и объединённая до колонки значения."""
    if value_col_idx > 1:
        ws.merge_cells(start_row=row, end_row=row, start_column=1, end_column=value_col_idx - 1)
    label_cell = ws.cell(row=row, column=1)
    label_cell.value = label
    label_cell.alignment = Alignment(horizontal="right")
    label_cell.font = Font(bold=True)
    value_cell = ws.cell(row=row, column=value_col_idx)
    value_cell.alignment = Alignment(horizontal="center")
    value_cell.font = Font(bold=True)


def _footer_cells(row: dict) -> list[dict]:
    """Итоговые ячейки строки footer_rows; поддерживает старый формат value_column + formula."""
    if "cells" in row:
        cells = row["cells"]
        return cells if isinstance(cells, list) else []
    if row.get("value_column"):
        return [{"column": row["value_column"], "formula": row.get("formula")}]
    return []


def _write_footer_label(ws, row: int, label: str, label_idx: int, first_value_idx: int) -> None:
    """Подпись итоговой строки: объединяется от колонки подписи до первой итоговой ячейки."""
    if first_value_idx - 1 > label_idx:
        ws.merge_cells(start_row=row, end_row=row, start_column=label_idx, end_column=first_value_idx - 1)
    label_cell = ws.cell(row=row, column=label_idx)
    label_cell.value = label
    label_cell.alignment = Alignment(horizontal="right")
    label_cell.font = Font(bold=True)


def dataframe_to_excel_bytes(
    df: pd.DataFrame,
    sheet_name: str | None = None,
    template: dict | None = None,
) -> bytes:
    """
    Сериализует DataFrame в xlsx по шаблону и возвращает байты
    (для st.download_button).
    """
    template = template or EXCEL_TEMPLATE
    validate_template(template)

    columns = template["columns"]
    keys = [c["key"] for c in columns]
    top_rows = template.get("top_rows", [])
    footer_rows = template.get("footer_rows", [])
    col_letters = {key: get_column_letter(i) for i, key in enumerate(keys, start=1)}
    col_index = {key: i for i, key in enumerate(keys, start=1)}

    # Колонки, которых нет в df (например, добавили в шаблон после создания df), — пустые.
    df = df.reindex(columns=keys)
    df = df.reset_index(drop=True)

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or template.get("sheet_name") or "Лист1")[:31]

    # Раскладка: top_rows, затем 2 строки шапки, затем данные, затем footer_rows.
    header_row = len(top_rows) + 1
    data_start = header_row + 2
    data_end = data_start + len(df) - 1

    # Верхние служебные строки (ячейки ввода).
    top_cells: dict[str, str] = {}
    for r, top in enumerate(top_rows, start=1):
        value_idx = col_index[top["input_column"]]
        _write_label_row(ws, r, top["label"], value_idx)
        top_cells[top["id"]] = f"${get_column_letter(value_idx)}${r}"

    _write_header(ws, columns, header_row)

    # Данные.
    ranges = _find_item_ranges(df, columns)
    start_rows = {start for start, _ in ranges}

    for idx in range(len(df)):
        excel_row = data_start + idx
        for col_i, col in enumerate(columns, start=1):
            cell = ws.cell(row=excel_row, column=col_i)
            if col["source"] == "formula":
                # Формула — только в первой строке позиции, остальные ячейки объединены.
                if idx in start_rows or not _is_item_level(col):
                    cell.value = _render_formula(
                        col["formula"], col_letters, top_cells, excel_row, data_start, data_end
                    )
            else:
                cell.value = _to_excel_value(df.iloc[idx][col["key"]])
            if col.get("number_format"):
                cell.number_format = col["number_format"]
            cell.alignment = Alignment(wrap_text=bool(col.get("wrap")), vertical="top")

    _merge_item_cells(ws, columns, ranges, data_start)

    # Нижние служебные строки.
    for offset, footer in enumerate(footer_rows, start=1):
        footer_row = data_end + offset
        cells = _footer_cells(footer)
        label_idx = col_index[footer.get("label_column", keys[0])]
        first_value_idx = min(col_index[c["column"]] for c in cells)
        _write_footer_label(ws, footer_row, footer["label"], label_idx, first_value_idx)
        for cell_def in cells:
            value_idx = col_index[cell_def["column"]]
            cell = ws.cell(row=footer_row, column=value_idx)
            cell.value = _render_formula(
                cell_def["formula"], col_letters, top_cells, None, data_start, data_end
            )
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(bold=True)
            if cell_def.get("number_format"):
                cell.number_format = cell_def["number_format"]

    # Ширина колонок.
    for col_i, col in enumerate(columns, start=1):
        if col.get("width"):
            width = col["width"]
        else:
            lengths = [len(str(col["label"]))]
            for value in df[col["key"]].tolist():
                if not _is_empty(value):
                    lengths.append(max(len(line) for line in str(value).split("\n")))
            width = min(max(lengths) + 2, 60)
        ws.column_dimensions[get_column_letter(col_i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
