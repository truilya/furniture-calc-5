# utils/parsers.py
"""
Модуль отвечает ТОЛЬКО за извлечение текста из загруженных файлов.
Никакой бизнес-логики (обращения к ИИ, генерации Excel) — единственная ответственность.
"""
from __future__ import annotations
import io

import streamlit as st
from pypdf import PdfReader
from docx import Document


class ParsingError(Exception):
    """Ошибка извлечения текста из файла."""


def _read_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages_text).strip()


def _read_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts).strip()


@st.cache_data(show_spinner=False)
def extract_text_from_file(file_bytes: bytes, file_name: str) -> str:
    """
    Кэшируется по хешу байтов файла: при rerun (например, смена модели
    в селекторе) повторный парсинг не выполняется.
    """
    suffix = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    try:
        if suffix == "pdf":
            text = _read_pdf(file_bytes)
        elif suffix == "docx":
            text = _read_docx(file_bytes)
        else:
            raise ParsingError(f"Неподдерживаемый формат: .{suffix}")
    except ParsingError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParsingError(f"Не удалось прочитать '{file_name}': {exc}") from exc

    if not text:
        raise ParsingError(
            f"Из файла '{file_name}' не удалось извлечь текст "
            "(возможно, это скан без текстового слоя)."
        )
    return text